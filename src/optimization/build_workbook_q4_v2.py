# -*- coding: utf-8 -*-
"""Openpyxl workbook builders for the question-4 v2 results.

result4-2.xlsx: 计划购电量 (issue rows g), 充放电量 (execution blocks),
                紧急购电量 (execution intervals).
result4-3.xlsx: 计划购电量 (q0), 调整购电量 (qA), 充放电量, 紧急购电量.

Row conventions (§2.3 of the optimized plan): the workbook row for day d
holds the 144 ISSUE columns (00:10..23:50 + 0:00-0:10+1); the row total
equals the sum of its own 144 columns; the row cost uses attachment-4's
template-order prices.  12/31's last column is generated causally (never
mechanically zeroed).

Run:  python -m src.optimization.build_workbook_q4_v2
"""
from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from src.optimization.question2 import ROOT, T
from src.optimization.q4_v2_common import battery_blocks, emergency_rows

Q4_2_DIR = ROOT / "outputs/question4/v2/result4-2"
Q4_3_DIR = ROOT / "outputs/question4/v2/result4-3"


def serial(date) -> float:
    return (pd.Timestamp(date) - pd.Timestamp("1899-12-30")).total_seconds() / 86400.0


def template_prices(root: Path = ROOT) -> np.ndarray:
    raw = pd.read_excel(root / "附件/附件4.xlsx", sheet_name=0)
    return raw.iloc[:, 1:145].to_numpy(float)  # (365, 144) template order


def issue_grid(issue: pd.DataFrame, date) -> np.ndarray:
    row = issue[issue.date == pd.Timestamp(date)]
    return row[[f"g{j}" for j in range(T)]].to_numpy(float).ravel()


def _fill_plan_sheet_dated(ws, rows, dates):
    for i, (date, grid, total, cost) in enumerate(rows):
        r = i + 2
        ws.cell(r, 1).value = serial(date)
        for j in range(T):
            ws.cell(r, 2 + j).value = float(grid[j])
        ws.cell(r, 146).value = float(total)
        ws.cell(r, 147).value = float(cost)


def build_result4_2(result_dir: Path, template: Path, p_tpl: np.ndarray) -> dict:
    wb = openpyxl.load_workbook(template)
    issue = pd.read_csv(result_dir / "question2_issue.csv", parse_dates=["date"])
    frame = pd.read_csv(result_dir / "question2_schedule.csv", parse_dates=["date"])
    p_sheet = wb["计划购电量"]
    rows = []
    for _, row in issue.iterrows():
        date = pd.Timestamp(row["date"])
        if date < pd.Timestamp("2025-02-01") or date > pd.Timestamp("2025-12-31"):
            continue
        grid = row[[f"g{j}" for j in range(T)]].to_numpy(float)
        d_idx = (date - pd.Timestamp("2025-01-01")).days
        rows.append((date, grid, float(grid.sum()),
                     float(np.dot(p_tpl[d_idx], grid))))
    _fill_plan_sheet_dated(p_sheet, rows, [r[0] for r in rows])
    b_sheet = wb["充放电量"]
    blocks = battery_blocks(frame[frame.date >= pd.Timestamp("2025-02-01")])
    for i, block in enumerate(blocks):
        r = i + 2
        b_sheet.cell(r, 1).value = serial(block[0]) if block[0] else None
        b_sheet.cell(r, 2).value = block[1]
        b_sheet.cell(r, 3).value = float(block[2])
        b_sheet.cell(r, 4).value = float(block[3])
        b_sheet.cell(r, 5).value = block[4]
        b_sheet.cell(r, 6).value = float(block[5]) if block[5] is not None else None
    e_sheet = wb["紧急购电量"]
    em = emergency_rows(frame[frame.date >= pd.Timestamp("2025-02-01")])
    for i, row in enumerate(em):
        r = i + 2
        e_sheet.cell(r, 1).value = serial(row[0])
        e_sheet.cell(r, 2).value = row[1]
        e_sheet.cell(r, 3).value = float(row[2])
    out_path = result_dir / "result4-2.xlsx"
    wb.save(out_path)
    return {"output": str(out_path), "plan_rows": len(rows),
            "battery_rows": len(blocks), "emergency_rows": len(em)}


def build_result4_3(result_dir: Path, template: Path, p_tpl: np.ndarray) -> dict:
    wb = openpyxl.load_workbook(template)
    q0 = pd.read_csv(result_dir / MAIN_Q0_CSV, parse_dates=["date"])
    qA = pd.read_csv(result_dir / MAIN_QA_CSV, parse_dates=["date"])
    frame = pd.read_csv(result_dir / MAIN_SCHEDULE_CSV, parse_dates=["date"])
    rows0, rowsA = [], []
    for _, (r0, rA) in enumerate(zip(q0.iterrows(), qA.iterrows())):
        date = pd.Timestamp(r0[1]["date"])
        if date < pd.Timestamp("2025-02-01") or date > pd.Timestamp("2025-12-31"):
            continue
        g0 = r0[1][[f"g{j}" for j in range(T)]].to_numpy(float)
        gA = rA[1][[f"g{j}" for j in range(T)]].to_numpy(float)
        d_idx = (date - pd.Timestamp("2025-01-01")).days
        p = p_tpl[d_idx]
        face = float(np.dot(p, g0))
        settlement = float(np.dot(p, np.minimum(g0, gA)
                                  + 0.5 * np.maximum(g0 - gA, 0.0)
                                  + 1.5 * np.maximum(gA - g0, 0.0)))
        rows0.append((date, g0, float(g0.sum()), face))
        rowsA.append((date, gA, float(gA.sum()), settlement))
    _fill_plan_sheet_dated(wb["计划购电量"], rows0, [r[0] for r in rows0])
    _fill_plan_sheet_dated(wb["调整购电量"], rowsA, [r[0] for r in rowsA])
    b_sheet = wb["充放电量"]
    blocks = battery_blocks(frame[frame.date >= pd.Timestamp("2025-02-01")])
    for i, block in enumerate(blocks):
        r = i + 2
        b_sheet.cell(r, 1).value = serial(block[0]) if block[0] else None
        b_sheet.cell(r, 2).value = block[1]
        b_sheet.cell(r, 3).value = float(block[2])
        b_sheet.cell(r, 4).value = float(block[3])
        b_sheet.cell(r, 5).value = block[4]
        b_sheet.cell(r, 6).value = float(block[5]) if block[5] is not None else None
    e_sheet = wb["紧急购电量"]
    em = emergency_rows(frame[frame.date >= pd.Timestamp("2025-02-01")])
    for i, row in enumerate(em):
        r = i + 2
        e_sheet.cell(r, 1).value = serial(row[0])
        e_sheet.cell(r, 2).value = row[1]
        e_sheet.cell(r, 3).value = float(row[2])
    out_path = result_dir / "result4-3.xlsx"
    wb.save(out_path)
    return {"output": str(out_path), "plan_rows": len(rows0),
            "battery_rows": len(blocks), "emergency_rows": len(em)}


MAIN_Q0_CSV = "M612/question3_issue_q0.csv"
MAIN_QA_CSV = "M612/question3_issue_qA.csv"
MAIN_SCHEDULE_CSV = "M612/question3_schedule.csv"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--q4-2-dir", type=Path, default=Q4_2_DIR)
    parser.add_argument("--q4-3-dir", type=Path, default=Q4_3_DIR)
    args = parser.parse_args()
    p_tpl = template_prices(ROOT)
    r2 = build_result4_2(args.q4_2_dir, ROOT / "附件/附件5/result4-2.xlsx", p_tpl)
    r3 = build_result4_3(args.q4_3_dir, ROOT / "附件/附件5/result4-3.xlsx", p_tpl)
    print(r2)
    print(r3)


if __name__ == "__main__":
    main()
